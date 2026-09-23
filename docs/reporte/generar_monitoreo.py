# -*- coding: utf-8 -*-
"""Genera el documento del caso de estudio de monitoreo en formato Word.

Uso, desde la raíz del repositorio:

    python docs/reporte/generar_monitoreo.py

La fuente canónica del contenido es ``docs/monitoreo-metricas-y-alertas.md``; este guion
produce la versión entregable en Word. Si ambos discrepan, manda el markdown.

Lee ``docs/reporte/enlaces-monitoreo.json`` (issues, PR, ramas y mediciones). Lo que falte
se imprime como [PENDIENTE] resaltado en amarillo, nunca se inventa: un documento de
monitoreo que afirma cifras que nadie midió es exactamente el problema que este trabajo
existe para resolver.

Las capturas no se incrustan: quedan como recuadros amarillos para pegarlas a mano.

La salida cae en `.github/DRAFTS/`, que git ignora. Es deliberado: lo que se versiona es
la fuente y este generador, de modo que el documento siempre pueda reproducirse.
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
SALIDA = RAIZ / ".github" / "DRAFTS" / "Monitoreo_Metricas_SLO_y_Alertas.docx"
REPO = "https://github.com/vidanj/Gestion-Servicios-Profesionales"
PENDIENTE = "[PENDIENTE]"

_ruta_datos = AQUI / "enlaces-monitoreo.json"
DATOS = json.loads(_ruta_datos.read_text(encoding="utf-8-sig")) if _ruta_datos.exists() else {}
ITEMS = DATOS.get("items", {})
EJECUCIONES = DATOS.get("ejecuciones", {})
MEDICIONES = DATOS.get("mediciones", {})
AUTOR = DATOS.get("autor", "Efrén Campuzano")
CUENTA = DATOS.get("cuenta", "vidanj")
FECHA = DATOS.get("fecha", "23 de septiembre de 2026")


def dato(clave: str, campo_: str) -> str:
    valor = ITEMS.get(clave, {}).get(campo_)
    return PENDIENTE if valor in (None, "") else str(valor)


def medicion(clave: str) -> str:
    valor = MEDICIONES.get(clave)
    return PENDIENTE if valor in (None, "") else str(valor)


# ---------------------------------------------------------------- utilidades de formato


def hipervinculo(parrafo_, url: str, texto: str) -> None:
    if not url or url == PENDIENTE:
        run = parrafo_.add_run(texto if texto != PENDIENTE else "[PENDIENTE: enlace]")
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return
    r_id = parrafo_.part.relate_to(url, RT.HYPERLINK, is_external=True)
    enlace = OxmlElement("w:hyperlink")
    enlace.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    subrayado = OxmlElement("w:u")
    subrayado.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(subrayado)
    run.append(rpr)
    texto_xml = OxmlElement("w:t")
    texto_xml.text = texto
    texto_xml.set(qn("xml:space"), "preserve")
    run.append(texto_xml)
    enlace.append(run)
    parrafo_._p.append(enlace)


def campo(run, instruccion: str, texto_previo: str = "") -> None:
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruccion
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    previo = OxmlElement("w:t")
    previo.text = texto_previo
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instr, separador, previo, fin):
        run._r.append(elemento)


def parrafo(doc, texto: str, estilo: str | None = None):
    """Párrafo con soporte mínimo de **negritas** y `código`."""
    p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
    for i, trozo in enumerate(texto.split("**")):
        for j, sub in enumerate(trozo.split("`")):
            if not sub:
                continue
            run = p.add_run(sub)
            run.bold = i % 2 == 1
            if j % 2 == 1:
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
    return p


def vinetas(doc, elementos: list[str], numeradas: bool = False) -> None:
    estilo = "List Number" if numeradas else "List Bullet"
    for elemento in elementos:
        parrafo(doc, elemento, estilo)


def _estilo_tabla(t) -> None:
    try:
        t.style = "Light Grid Accent 1"
    except KeyError:
        t.style = "Table Grid"


def tabla(doc, encabezados: list[str], filas: list[list], anchos_cm: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    _estilo_tabla(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, encabezado in enumerate(encabezados):
        celda = t.rows[0].cells[i]
        celda.text = ""
        run = celda.paragraphs[0].add_run(encabezado)
        run.bold = True
        run.font.size = Pt(9)
    for fila in filas:
        celdas = t.add_row().cells
        for i, valor in enumerate(fila):
            celdas[i].text = ""
            p = celdas[i].paragraphs[0]
            if isinstance(valor, tuple):  # (texto, url) → hipervínculo
                hipervinculo(p, valor[1], valor[0])
            else:
                run = p.add_run(str(valor))
                if str(valor).startswith("[PENDIENTE"):
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            for run in p.runs:
                run.font.size = Pt(9)
    if anchos_cm:
        for fila in t.rows:
            for i, ancho in enumerate(anchos_cm):
                fila.cells[i].width = Cm(ancho)
    doc.add_paragraph()
    return t


def codigo(doc, texto: str) -> None:
    for linea in texto.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        sombreado = OxmlElement("w:shd")
        sombreado.set(qn("w:val"), "clear")
        sombreado.set(qn("w:fill"), "F2F2F2")
        p._p.get_or_add_pPr().append(sombreado)
        run = p.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
    doc.add_paragraph()


def marcador(doc, texto: str, alto_cm: float = 6.0) -> None:
    """Recuadro vacío con la instrucción de qué pegar."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fila = t.rows[0]
    fila.height = Cm(alto_cm)
    fila.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    p = fila.cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"[{texto}]")
    run.bold = True
    run.font.size = Pt(10)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph()


def evidencia(doc, clave: str, punto: str, capturas: list[str] | None = None) -> None:
    """Bloque estándar de evidencia: issue, PR, rama, commit y recuadros de captura."""
    doc.add_heading(f"Evidencia — {punto}", level=2)
    parrafo(doc, f"**{ITEMS.get(clave, {}).get('titulo', PENDIENTE)}**")
    p = doc.add_paragraph()
    p.add_run("Issue: ").bold = True
    hipervinculo(p, dato(clave, "issue_url"), f"#{dato(clave, 'issue')}")
    p.add_run("   ·   PR: ").bold = True
    hipervinculo(p, dato(clave, "pr_url"), f"#{dato(clave, 'pr')}")
    p2 = doc.add_paragraph()
    p2.add_run("Rama: ").bold = True
    p2.add_run(dato(clave, "rama")).font.name = "Consolas"
    p2.add_run("   ·   Commit: ").bold = True
    p2.add_run(dato(clave, "commit")).font.name = "Consolas"
    p2.add_run("   ·   Hecho por: ").bold = True
    p2.add_run(f"{AUTOR} (@{CUENTA})")
    for texto in capturas or [
        f"CAPTURA: issue #{dato(clave, 'issue')} en GitHub, con sus etiquetas y su descripción",
        f"CAPTURA: PR #{dato(clave, 'pr')} — pestaña de conversación",
        f"CAPTURA: PR #{dato(clave, 'pr')} — pestaña de archivos, con el resumen de cambios",
    ]:
        marcador(doc, texto, 5.5)


def ejecucion(doc, clave: str, titulo: str, descripcion: str) -> None:
    """Bloque de evidencia de una ejecución de integración continua."""
    doc.add_heading(titulo, level=3)
    parrafo(doc, descripcion)
    url = EJECUCIONES.get(clave) or ""
    p = doc.add_paragraph()
    p.add_run("Ejecución: ").bold = True
    hipervinculo(p, url, url if url else PENDIENTE)
    marcador(doc, f"CAPTURA: {titulo} — pasos en verde", 5.0)


def nota(doc, texto: str) -> None:
    """Aviso destacado en una caja de una sola celda."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    p = t.rows[0].cells[0].paragraphs[0]
    for i, trozo in enumerate(texto.split("**")):
        run = p.add_run(trozo)
        run.bold = i % 2 == 1
        run.font.size = Pt(9.5)
    doc.add_paragraph()


# ---------------------------------------------------------------- documento


def configurar(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for seccion in doc.sections:
        seccion.left_margin = seccion.right_margin = Cm(2.2)
        seccion.top_margin = seccion.bottom_margin = Cm(2.0)
        pie = seccion.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pie.add_run("Monitoreo de la aplicación — Gestión de Servicios Profesionales — pág. ")
        campo(pie.add_run(), "PAGE", "1")
    actualizar = OxmlElement("w:updateFields")
    actualizar.set(qn("w:val"), "true")
    doc.settings.element.append(actualizar)


def portada(doc) -> None:
    for _ in range(3):
        doc.add_paragraph()
    for texto, tam, negrita in (
        ("[INSTITUCIÓN]", 14, True),
        ("[MATERIA] — [GRUPO]", 12, False),
        ("", 12, False),
        ("Monitoreo de la aplicación", 26, True),
        ("Métricas, niveles de servicio y alertas del pipeline de liberación", 16, False),
        ("Sistema de Gestión de Servicios Profesionales", 14, True),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(texto)
        run.bold = negrita
        run.font.size = Pt(tam)
        if texto.startswith("["):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph()
    tabla(
        doc,
        ["Campo", "Valor"],
        [
            ["Docente", "[NOMBRE DEL DOCENTE]"],
            ["Equipo", f"{AUTOR} · [agregar a los demás integrantes]"],
            ["Repositorio", (REPO, REPO)],
            ["Rama de integración", "dev"],
            ["Especificación", "Spec 008 — monitoreo, niveles de servicio y despliegue verificado"],
            ["Fecha", FECHA],
            ["Alcance", "Punto 1 de 3: métricas para el monitoreo"],
        ],
        [4, 11],
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def indice(doc) -> None:
    doc.add_heading("Índice", level=1)
    campo(
        doc.add_paragraph().add_run(),
        'TOC \\o "1-3" \\h \\z \\u',
        "Clic derecho → Actualizar campo para generar el índice.",
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def alcance(doc) -> None:
    doc.add_heading("Objetivo y alcance", level=1)
    parrafo(
        doc,
        "Este documento describe y justifica el **flujo de trabajo de liberación y despliegue "
        "continuo** de la aplicación, el **entorno requerido**, los **niveles de servicio "
        "acordados**, las **métricas establecidas para el monitoreo** y los **parámetros de "
        "configuración de las herramientas** utilizadas.",
    )
    doc.add_heading("Los tres puntos del caso de estudio", level=2)
    tabla(
        doc,
        ["Punto", "Alcance", "Estado"],
        [
            ["1. Métricas para el monitoreo",
             "Recolección, tableros, niveles de servicio, alarmas y despliegue verificado",
             "Desarrollado aquí"],
            ["2. Visor de trazabilidad", "Registros y trazas", "Reservado — sección 13"],
            ["3. Visor de auditoría", "Bitácora de acciones de usuario", "Reservado — sección 14"],
        ],
        [5.0, 7.0, 3.5],
    )
    doc.add_heading("Qué existe ya y qué es diseño aprobado", level=2)
    parrafo(
        doc,
        "La distinción se mantiene en todo el documento. **Lo que no está medido se dice que no "
        "lo está**: un informe de monitoreo que afirma cifras que nadie tomó es precisamente el "
        "problema que este trabajo existe para resolver.",
    )
    tabla(
        doc,
        ["Componente", "Estado"],
        [
            ["Registros estructurados con correlación por identificador de traza",
             "En producción (issue #124)"],
            ["Instrumentación de métricas y trazas", "En producción, sin consumidor (issue #124)"],
            ["Sondas de vida y de preparación", "En producción (issue #123)"],
            ["Verificación de la imagen en integración continua",
             "En producción (issue #131, primera mitad)"],
            ["Entorno de monitoreo, tableros y alarmas",
             f"Diseñado y aprobado, pendiente de construir (issue #{dato('implementacion', 'issue')})"],
            ["Despliegue automatizado con verificación y reversión",
             f"Diseñado y aprobado, pendiente de construir (issues #{dato('implementacion', 'issue')} y #189)"],
        ],
        [9.5, 6.0],
    )


def seccion_1(doc) -> None:
    doc.add_heading("1. Por qué esto, y por qué ahora", level=1)
    parrafo(
        doc,
        "El sistema ya produce telemetría que nadie consume. La configuración de telemetría del "
        "backend registra las métricas de la plataforma web y las exporta **solo si** está "
        "definida la variable que apunta al recolector. Esa variable está vacía desde que se "
        "escribió el código, porque no existía recolector alguno al que apuntar. El resultado es "
        "que hoy la única manera de saber si la aplicación funciona es abrirla en un navegador.",
    )
    parrafo(
        doc,
        "En paralelo, el despliegue se lanza a mano desde el panel del proveedor, la versión en "
        "marcha no se puede ligar a un commit y no hay procedimiento de reversión. Los objetivos "
        "de entrega que el proyecto se fijó en su estrategia de despliegue dicen literalmente "
        "«desconocido» y «no se mide».",
    )
    nota(
        doc,
        "Son la misma carencia vista desde dos lados, y por eso se resuelven juntas: **un "
        "despliegue sin verificación es una apuesta, y un monitoreo sin despliegue automatizado "
        "solo sirve para enterarse tarde.**",
    )


def seccion_2(doc) -> None:
    doc.add_heading("2. Decisión estructural: dónde viven las métricas", level=1)
    parrafo(
        doc,
        "La aplicación **no expone un endpoint de métricas**. Empuja su telemetría por OTLP a un "
        "colector de OpenTelemetry, y es el colector quien la publica en el formato que el "
        "recolector entiende.",
    )
    marcador(doc, "DIAGRAMA: API → colector → recolector → tableros y alarmas", 5.5)
    doc.add_heading("Las tres razones, en orden de peso", level=2)
    parrafo(doc, "**1. El exportador que haría falta no tiene versión estable.** Consultado el 23 "
                 "de septiembre de 2026 en el repositorio oficial de paquetes:")
    tabla(
        doc,
        ["Paquete", "Última versión", "Última estable"],
        [
            ["OpenTelemetry.Exporter.Prometheus.AspNetCore", "1.19.1-beta.1", "ninguna"],
            ["OpenTelemetry.Instrumentation.EntityFrameworkCore", "1.19.0-beta.1", "ninguna"],
            ["OpenTelemetry.Exporter.OpenTelemetryProtocol", "1.19.1", "1.19.1 (ya instalado)"],
        ],
        [8.0, 3.8, 3.7],
    )
    parrafo(
        doc,
        "El proyecto compila tratando las advertencias como errores y prohíbe silenciarlas, así "
        "que ninguno de los dos primeros cabe. Esto **corrige una suposición del issue #187**, "
        "que aplazó las métricas «hasta que los paquetes salgan de preestreno» dando por hecho "
        "que era cuestión de tiempo. Año y medio después siguen ahí, y ninguno ha publicado jamás "
        "una versión estable.",
    )
    parrafo(
        doc,
        "**2. Un endpoint de métricas en una API pública es superficie de información**: revela "
        "rutas, volúmenes de uso y versiones. Protegerlo exigiría autenticación propia o filtrado "
        "por red, que el proveedor actual no ofrece.",
    )
    parrafo(
        doc,
        "**3. El colector desacopla.** La aplicación empuja y no le importa quién consume. "
        "Sustituir el recolector por otro sistema no toca una línea de código.",
    )
    nota(
        doc,
        "La decisión se protege con una **prueba de integración** que afirma que la ruta de "
        "métricas responde «no encontrado». Sin ella, alguien la añadiría más adelante sin "
        "advertir que contradice esto.",
    )


def seccion_3(doc) -> None:
    doc.add_heading("3. Estado de partida", level=1)
    tabla(
        doc,
        ["Aspecto", "Antes de esta entrega"],
        [
            ["Registros", "Estructurados en JSON, con identificador de traza y redacción de secretos"],
            ["Trazas", "Instrumentadas (web, cliente HTTP y base de datos), sin exportar"],
            ["Métricas", "Instrumentadas, sin consumidor"],
            ["Tableros", "No existen"],
            ["Alarmas", "No existen"],
            ["Niveles de servicio", "Solo un umbral de latencia, en las pruebas de carga"],
            ["Despliegue", "Manual, desde el panel del proveedor"],
            ["Reversión", "Sin procedimiento"],
            ["Trazabilidad de la versión desplegada", "Inexistente"],
        ],
        [5.5, 10.0],
    )


def seccion_4(doc) -> None:
    doc.add_heading("4. Flujo de trabajo del pipeline", level=1)
    doc.add_heading("4.1 El recorrido completo", level=2)
    marcador(doc, "DIAGRAMA: del commit a la versión verificada, con sus puertas", 9.0)
    doc.add_heading("4.2 Etapas", level=2)
    tabla(
        doc,
        ["Etapa", "Herramienta", "Disparador", "Puerta", "Estado"],
        [
            ["Formato y análisis", "Actions + CSharpier, StyleCop, Roslynator", "PR y push",
             "Cero advertencias", "En uso"],
            ["Pruebas", "Actions + xUnit, coverlet", "PR y push", "Suite en verde", "En uso"],
            ["Frontend", "Actions + Playwright", "PR con cambios", "Compilación y E2E", "En uso"],
            ["Imagen", "Actions + Docker Buildx", "PR y push", "Arranca y responde sano", "En uso"],
            ["Seguridad", "CodeQL + Dependabot", "PR y semanal", "Sin alertas nuevas", "En uso"],
            ["Pruebas de carga", "Actions + k6", "Manual o por etiqueta", "p95 bajo umbral", "En uso"],
            ["Análisis estático", "SonarQube local", "Manual", "Puerta sobre código nuevo", "En uso"],
            ["Liberación", "Actions + disparador del proveedor", "Imagen verificada en dev",
             "Las cuatro comprobaciones", "Por construir"],
            ["Despliegue", "Proveedor de alojamiento", "Tras la aprobación", "—", "Por construir"],
            ["Verificación", "Sondeo de la sonda de preparación", "Tras desplegar",
             "Tres respuestas sanas", "Por construir"],
            ["Reversión", "El mismo flujo, commit anterior", "Manual", "La misma verificación",
             "Por construir"],
        ],
        [3.2, 4.2, 3.0, 3.2, 2.0],
    )
    doc.add_heading("4.3 Por qué cada puerta está donde está", level=2)
    for titulo, cuerpo in (
        ("Por qué el despliegue exige las cuatro comprobaciones y no solo la última",
         "Encadenar el despliegue al flujo que construye la imagen demostraría únicamente que la "
         "imagen se construye. Un commit con las pruebas en rojo pasaría. Por eso la puerta "
         "consulta el estado de las cuatro comprobaciones **para ese commit concreto**, y no el "
         "resultado del flujo que la invocó."),
        ("Por qué se sondea la sonda de preparación y no la de vida",
         "La de vida responde en cuanto el proceso arranca. La de preparación, además, confirma "
         "que la base de datos es alcanzable y, por tanto, que las migraciones se aplicaron. Es la "
         "diferencia entre «el contenedor está encendido» y «el sistema sirve»."),
        ("Por qué hay una espera antes del primer sondeo",
         "Sin ella se sondearía la instancia **anterior**, que responde perfectamente sana, y el "
         "despliegue se daría por bueno sin haber ocurrido. Es el modo de fallo más probable de "
         "todo el diseño y el más difícil de detectar, porque produce un pipeline verde."),
        ("Por qué se exigen tres respuestas sanas consecutivas",
         "Durante un reinicio el servicio rebota: puede responder sano un instante y caer de "
         "nuevo. Una sola respuesta no distingue «arrancó» de «está arrancando»."),
        ("Por qué la concurrencia no cancela",
         "Dos despliegues simultáneos pueden ejecutar migraciones a la vez. Se serializan por "
         "construcción, no por suerte."),
        ("Por qué el plazo de verificación es largo",
         "El plan contratado duerme la instancia por inactividad y además construye la imagen en "
         "su propia infraestructura. Un plazo corto reportaría como fallo lo que es un arranque "
         "normal, y la primera consecuencia sería que alguien lo ampliara sin entender por qué, o "
         "peor, que desactivara la verificación."),
        ("Por qué la reversión no es un flujo aparte",
         "Revertir es desplegar un commit anterior con la misma verificación. Así el camino de "
         "reversión **se ejercita en cada despliegue**, en lugar de ser una ruta que nadie ha "
         "recorrido hasta el día del incidente."),
    ):
        doc.add_heading(titulo, level=3)
        parrafo(doc, cuerpo)
    doc.add_heading("4.4 Límite conocido: qué significa «trazable» aquí", level=2)
    parrafo(
        doc,
        "El disparador del proveedor hace que la plataforma **reconstruya** la imagen a partir "
        "del commit. La imagen que llega a producción no es, por tanto, el mismo artefacto binario "
        "que se verificó en integración continua, sino uno equivalente construido en otro sitio.",
    )
    nota(
        doc,
        "La trazabilidad que se obtiene es **commit → despliegue**, no **artefacto → despliegue**. "
        "Queda escrito para que nadie lea «imagen inmutable por commit» en la estrategia de "
        "despliegue y suponga algo más fuerte de lo que hay.",
    )


def seccion_5(doc) -> None:
    doc.add_heading("5. Entorno requerido", level=1)
    doc.add_heading("5.1 Entorno de ejecución de la aplicación", level=2)
    tabla(
        doc,
        ["Elemento", "Valor"],
        [
            ["Alojamiento", "Proveedor de contenedores, plan gratuito, detrás de una red de distribución"],
            ["Imagen", "Multietapa: compilación, paquete de migraciones y ejecución"],
            ["Puerto", "10000"],
            ["Base de datos", "PostgreSQL 18"],
            ["Migraciones", "Paquete autocontenido ejecutado al arrancar"],
            ["Sonda del contenedor", "Preparación, cada 30 s, con 60 s de gracia inicial"],
            ["Características del plan que condicionan el diseño",
             "Memoria acotada y suspensión por inactividad, con arranques en frío de decenas de segundos"],
        ],
        [5.5, 10.0],
    )
    doc.add_heading("5.2 Entorno de monitoreo", level=2)
    parrafo(
        doc,
        "Vive en su propio archivo de composición, **independiente del de la aplicación**. La "
        "razón es el ciclo de vida: el monitoreo tiene que seguir en pie precisamente cuando la "
        "aplicación se cae, y si compartieran proyecto, derribar la aplicación se llevaría por "
        "delante la evidencia del incidente.",
    )
    tabla(
        doc,
        ["Servicio", "Función", "Puerto", "Estado persistente"],
        [
            ["Colector de OpenTelemetry", "Recibe OTLP y republica", "4317, 4318, 8889", "No"],
            ["Recolector de métricas", "Almacena, evalúa reglas, dispara alarmas", "9090", "Volumen"],
            ["Gestor de alertas", "Agrupa, silencia, inhibe y entrega", "9093", "Volumen"],
            ["Tableros", "Visualización", "3001", "Volumen"],
            ["Sondeo externo", "Observa las sondas desde fuera", "9115", "No"],
        ],
        [4.5, 5.5, 3.0, 2.5],
    )
    nota(
        doc,
        "**El tablero se publica en 3001 y no en 3000** porque el frontend ocupa ese puerto. Es el "
        "tipo de choque que cuesta media hora de diagnóstico la primera vez.",
    )
    doc.add_heading("5.3 Los dos alcances de observación", level=2)
    parrafo(doc, "Es la limitación más importante de este diseño y conviene entenderla antes de "
                 "leer cualquier cifra.")
    tabla(
        doc,
        ["Alcance", "Qué se observa", "Cómo"],
        [
            ["Entorno local", "Todo: latencia por ruta, memoria, recolección de basura, cola de "
                              "hilos y métricas de negocio", "La aplicación empuja telemetría"],
            ["Servicio publicado", "Solo disponibilidad y tiempo de respuesta desde fuera",
             "Sondeo externo de las sondas de salud"],
        ],
        [3.5, 7.5, 4.5],
    )
    parrafo(
        doc,
        "El contenedor publicado **no puede** empujar telemetría a un colector que corre en el "
        "equipo de una persona, y exponer ese colector a internet sería abrir un receptor público "
        "de telemetría. A cambio, el sondeo externo da exactamente los dos datos que hacen falta "
        "del entorno real: si está disponible y cuánto tarda en despertar.",
    )


def seccion_6(doc) -> None:
    doc.add_heading("6. Niveles de servicio acordados", level=1)
    nota(
        doc,
        "**El arranque en frío se mide y se declara, pero no consume presupuesto de error.** Es "
        "una característica conocida y aceptada del plan contratado, no un defecto del software. "
        "Confundir ambas cosas llevaría a perseguir un objetivo que solo se alcanza pagando otro "
        "plan.",
    )
    tabla(
        doc,
        ["ID", "Indicador", "Objetivo", "Ventana", "Consecuencia si se incumple"],
        [
            ["SLO-1", "Sondeos de preparación con estado sano, excluidos 90 s tras arranque en frío",
             "≥ 99,0 %", "30 días", "Incidencia; se evalúa un plan con instancia siempre activa"],
            ["SLO-2", "Sondeos de vida con estado sano", "≥ 99,5 %", "30 días",
             "Investigación de reinicios: memoria o fallo de arranque"],
            ["SLO-3", "Percentil 95 del tiempo de respuesta", "≤ 1,5 s interno · ≤ 5 s acordado",
             "7 días", "Sobre 1,5 s, tarea de rendimiento; sobre 5 s, bloquea el siguiente despliegue"],
            ["SLO-4", "Proporción de errores de servidor", "≤ 1 % (alarma al 5 %)", "7 días",
             "La corrección tiene prioridad sobre funcionalidad nueva"],
            ["SLO-5", "Duración del primer acceso tras inactividad", "≤ 90 s (p95)", "30 días",
             "Informativo: se declara para no confundirlo con una caída"],
            ["SLO-6", "Tiempo con métricas disponibles", "≥ 99 %", "7 días",
             "Sin métricas no hay ninguno de los objetivos anteriores"],
            ["SLO-7", "Respaldos que terminan correctamente", "100 %", "30 días",
             "Incidente crítico: afecta a la recuperabilidad"],
            ["SLO-8", "Despliegues sanos a la primera", "≥ 95 %", "30 días / 20 despliegues",
             "Se endurecen las puertas antes de seguir desplegando"],
            ["SLO-9", "De versión defectuosa a servicio sano", "≤ 10 min", "Por incidente",
             "Se revisa el procedimiento de reversión"],
            ["SLO-10", "De integrar a servicio sano verificado", "≤ 20 min", "30 días",
             "Se recorta o paraleliza la etapa más lenta"],
        ],
        [1.5, 4.6, 2.6, 2.2, 4.6],
    )
    doc.add_heading("De dónde salen estos umbrales", level=2)
    parrafo(
        doc,
        "**SLO-3 hereda su techo de las pruebas de carga.** El proyecto ya fijó 5 s en el "
        "percentil 95 como criterio de fallo de su plan de pruebas de carga. Declarar aquí un "
        "umbral distinto crearía dos verdades sobre la misma magnitud, y la primera vez que "
        "discreparan nadie sabría cuál rige. El objetivo interno más estricto existe porque 5 s "
        "ya es una experiencia mala: esperar a rozarlo para reaccionar es llegar tarde.",
    )
    parrafo(
        doc,
        "**SLO-8, SLO-9 y SLO-10** ya estaban comprometidos en la estrategia de despliegue del "
        "proyecto como métricas de entrega. Hasta ahora no tenían **de dónde** medirse; el flujo "
        "de despliegue verificado es esa fuente.",
    )
    parrafo(
        doc,
        "**El reloj de SLO-10 no incluye la espera de aprobación humana.** Si la incluyera, "
        "mediría la disponibilidad del revisor en lugar del pipeline.",
    )


def seccion_7(doc) -> None:
    doc.add_heading("7. Métricas establecidas para el monitoreo", level=1)
    doc.add_heading("7.1 Cómo se leen los nombres", level=2)
    parrafo(
        doc,
        "La traducción la hace el colector y es mecánica: los puntos pasan a guiones bajos, se "
        "añade el sufijo de la unidad y los contadores reciben un sufijo de total. Un histograma "
        "produce tres series: los tramos, la suma y el conteo.",
    )
    codigo(
        doc,
        "http.server.request.duration  ->  http_server_request_duration_seconds_bucket\n"
        "                                  http_server_request_duration_seconds_sum\n"
        "                                  http_server_request_duration_seconds_count",
    )
    doc.add_heading("7.2 Peticiones y servidor", level=2)
    parrafo(doc, "Ya se emiten hoy; lo único que faltaba era alguien que las consumiera.")
    tabla(
        doc,
        ["Métrica", "Tipo", "Para qué sirve"],
        [
            ["http_server_request_duration_seconds", "Histograma",
             "La métrica central: de ella salen latencia p95, proporción de errores y volumen"],
            ["http_server_active_requests", "Contador bidireccional", "Saturación instantánea"],
            ["kestrel_active_connections", "Contador bidireccional", "Conexiones abiertas"],
            ["kestrel_queued_requests", "Contador bidireccional", "Trabajo aceptado y no atendido"],
            ["kestrel_rejected_connections_total", "Contador", "Límite de conexiones alcanzado"],
            ["aspnetcore_routing_match_attempts_total", "Contador", "Tráfico a rutas inexistentes"],
            ["aspnetcore_diagnostics_exceptions_total", "Contador", "Excepciones no controladas"],
            ["aspnetcore_rate_limiting_requests_total", "Contador",
             "Rechazos del limitador de autenticación"],
        ],
        [6.5, 3.2, 5.8],
    )
    nota(
        doc,
        "La última exige registrar su medidor **explícitamente**: no viaja con la instrumentación "
        "general de la plataforma web, aunque parezca que debería. Sin esa línea no hay forma de "
        "ver cuántas peticiones rechaza el limitador, que es justo la señal que distingue «nadie "
        "entra» de «alguien está probando contraseñas».",
    )
    doc.add_heading("7.3 Tiempo de ejecución de la plataforma", level=2)
    tabla(
        doc,
        ["Métrica", "Para qué sirve"],
        [
            ["dotnet_process_memory_working_set_bytes",
             "La más importante del grupo: el plan tiene memoria acotada y el reinicio por memoria "
             "es la caída más frecuente"],
            ["dotnet_gc_collections_total",
             "Presión de memoria; el crecimiento sostenido de la generación mayor delata una fuga"],
            ["dotnet_gc_pause_time_seconds_total",
             "Tiempo con la aplicación detenida: explica latencias sin causa aparente"],
            ["dotnet_thread_pool_queue_length",
             "Trabajo pendiente: si crece, la latencia subirá a continuación"],
            ["dotnet_monitor_lock_contentions_total", "Contención de bloqueos"],
            ["dotnet_exceptions_total", "Excepciones lanzadas, incluidas las capturadas"],
        ],
        [6.5, 9.0],
    )
    nota(
        doc,
        "**Aviso sobre la documentación existente.** En la versión actual de la plataforma, el "
        "paquete de instrumentación ya no emite los nombres con prefijo de proceso: registra el "
        "medidor integrado, cuyos instrumentos tienen otro prefijo. Casi todos los tableros "
        "publicados que se encuentran buscando fueron escritos para versiones anteriores; "
        "copiarlos produce un tablero que se aprovisiona **sin ningún error** y aparece "
        "permanentemente vacío.",
    )
    doc.add_heading("7.4 Métricas de negocio", level=2)
    parrafo(doc, "Son las que distinguen «la API responde 200» de «el negocio funciona».")
    tabla(
        doc,
        ["Métrica", "Etiqueta y valores", "Para qué sirve"],
        [
            ["gsp_solicitudes_creadas_total", "resultado: creada, servicio_no_disponible",
             "Volumen de negocio. Una caída a cero con tráfico normal significa que el flujo se "
             "rompió donde el 200 no lo delata"],
            ["gsp_solicitudes_cambios_estado_total",
             "estado_origen, estado_destino, resultado: aceptada, rechazada, no_autorizada",
             "Embudo del negocio y detección de defectos"],
            ["gsp_autenticacion_intentos_total",
             "resultado: exito, credenciales_invalidas, cuenta_inactiva",
             "Fuerza bruta y degradación del acceso"],
            ["gsp_usuarios_registrados_total", "resultado: creado, correo_duplicado",
             "Crecimiento y detección de altas automatizadas"],
            ["gsp_respaldos_ejecutados_total",
             "resultado: exito, herramienta_ausente, fallo_ejecucion, configuracion_incompleta",
             "Convierte «hay guiones de respaldo» en «los respaldos se ejecutan»"],
            ["gsp_respaldos_duracion_seconds", "resultado",
             "El respaldo crece con la base; la tendencia avisa antes de que agote su plazo"],
        ],
        [5.0, 5.5, 5.0],
    )
    doc.add_heading("Dos decisiones que merecen explicación", level=3)
    parrafo(
        doc,
        "**Sobre el inicio de sesión.** El servicio devuelve deliberadamente el mismo mensaje para "
        "credenciales incorrectas y para cuenta desactivada, para no revelar cuál de las dos "
        "ocurrió. La métrica sí distingue ambos casos, y **eso no rompe esa decisión**: la "
        "respuesta que ve el cliente sigue siendo idéntica, y la métrica es agregada, interna y "
        "sin identificadores. Es precisamente lo que permite responder «¿los usuarios no entran "
        "porque se equivocan, o porque sus cuentas están desactivadas?» sin filtrar nada a quien "
        "pregunta desde fuera.",
    )
    parrafo(
        doc,
        "**Sobre los cambios de estado.** Las transiciones válidas del sistema son seis; dos "
        "estados son terminales. Por eso un volumen apreciable de transiciones rechazadas no "
        "significa tráfico hostil: significa que la interfaz está ofreciendo transiciones que el "
        "servidor no admite. Es un defecto de producto que ninguna métrica técnica mostraría.",
    )
    doc.add_heading("7.5 Sondeo externo", level=2)
    tabla(
        doc,
        ["Métrica", "Para qué sirve"],
        [
            ["probe_success", "Disponibilidad. Base de SLO-1 y SLO-2"],
            ["probe_duration_seconds",
             "Tiempo de respuesta desde fuera; en el servicio publicado mide el arranque en frío"],
            ["probe_http_status_code", "Distingue un 503 de la sonda de un error de red"],
        ],
        [5.0, 10.5],
    )
    parrafo(
        doc,
        "El sondeo no se conforma con el código 200: afirma también el contenido del cuerpo, "
        "porque un estado degradado podría responder 200 y pasar por sano.",
    )
    doc.add_heading("7.6 Política de etiquetado y cardinalidad", level=2)
    parrafo(
        doc,
        "Cada etiqueta multiplica el número de series almacenadas. Una etiqueta con valores "
        "ilimitados no degrada el sistema poco a poco: lo tumba.",
    )
    parrafo(doc, "**Prohibido como etiqueta:**")
    vinetas(doc, [
        "Correo, nombre y teléfono.",
        "Identificadores de usuario, solicitud o servicio.",
        "Direcciones IP.",
        "Rutas con parámetros ya sustituidos.",
        "Mensajes de excepción y marcas de tiempo.",
    ])
    nota(
        doc,
        "La regla operativa es una sola pregunta: **¿cuántos valores distintos puede tomar esta "
        "etiqueta a lo largo de un año?** Si la respuesta no es un número pequeño escrito en el "
        "catálogo, no es una etiqueta: es un registro. **La métrica cuenta; el registro "
        "identifica.**",
    )
    parrafo(
        doc,
        "Peor caso de la métrica con más etiquetas: 5 × 5 × 3 = 75 series, y en la práctica muchas "
        "menos, porque la mayoría de las combinaciones nunca ocurre.",
    )


def seccion_8(doc) -> None:
    doc.add_heading("8. Alarmas y alertas", level=1)
    doc.add_heading("8.1 Las reglas", level=2)
    parrafo(
        doc,
        "Diez reglas, todas derivadas de un nivel de servicio de la sección 6. Las expresiones "
        "compartidas con los tableros se definen **una sola vez** como reglas de registro, para "
        "que tablero y alarma no calculen lo mismo por separado y acaben discrepando.",
    )
    tabla(
        doc,
        ["#", "Alarma", "Condición", "Espera", "Severidad", "SLO"],
        [
            ["A1", "Servicio no disponible", "La sonda de preparación no responde sana", "3 min",
             "Crítica", "SLO-1"],
            ["A2", "Base de datos inalcanzable", "La sonda de vida responde y la de preparación no",
             "2 min", "Crítica", "SLO-1"],
            ["A3", "Sin métricas", "Las series de la aplicación desaparecieron", "10 min",
             "Advertencia", "SLO-6"],
            ["A4", "Tasa de error alta", "Más del 5 % de errores, con tráfico significativo",
             "5 min", "Crítica", "SLO-4"],
            ["A5", "Latencia degradada", "p95 sobre el objetivo interno", "10 min", "Advertencia",
             "SLO-3"],
            ["A6", "Latencia fuera del acuerdo", "p95 sobre el acuerdo formal", "5 min", "Crítica",
             "SLO-3"],
            ["A7", "Memoria cerca del límite", "Conjunto de trabajo sobre el umbral", "15 min",
             "Advertencia", "SLO-2"],
            ["A8", "Cola de trabajo creciente", "Trabajo pendiente acumulado", "5 min",
             "Advertencia", "SLO-3"],
            ["A9", "Autenticaciones fallidas", "Ritmo de fallos sobre el umbral", "10 min",
             "Advertencia", "—"],
            ["A10", "Respaldo fallido", "Al menos uno fallido en la última hora", "Inmediata",
             "Crítica", "SLO-7"],
        ],
        [1.0, 3.6, 5.2, 1.8, 2.2, 1.7],
    )
    doc.add_heading("8.2 Decisiones que no son obvias", level=2)
    for titulo, cuerpo in (
        ("A2 distingue «el proceso vive y la base no responde»",
         "Es posible porque la comprobación de base de datos es la única etiquetada como parte de "
         "la preparación: si la sonda de vida responde y la de preparación no, la causa es una "
         "sola. Sin esta regla, el diagnóstico se haría a mano en mitad del incidente."),
        ("A4 exige tráfico mínimo",
         "Con tráfico casi nulo, una sola respuesta de error da el 100 % de tasa de error. Sin esa "
         "condición, la alarma se dispararía cada noche."),
        ("Hay dos alarmas de latencia y no una",
         "A5 avisa con margen y severidad menor; A6 señala un incumplimiento del acuerdo. Una sola "
         "regla obligaría a elegir entre avisar tarde o avisar siempre."),
        ("Las esperas están calibradas para absorber el arranque en frío",
         "Si A1 disparase en un minuto, avisaría cada vez que el servicio despierta de una siesta. "
         "La consecuencia no sería un aviso de más: sería que en dos semanas **nadie se cree las "
         "alarmas**. Hay un caso de prueba que afirma explícitamente que un arranque en frío de 60 "
         "segundos no dispara nada."),
        ("Las derivadas se inhiben",
         "Una aplicación caída dispara además latencia, error y saturación. Sin inhibición "
         "llegarían cuatro avisos del mismo hecho y el ruido taparía la causa."),
        ("El destino externo se deja configurado y desactivado",
         "Un entorno de pruebas que envía correos en cada ensayo se vuelve inusable en una semana, "
         "y entonces alguien lo desactiva sin documentarlo."),
    ):
        doc.add_heading(titulo, level=3)
        parrafo(doc, cuerpo)
    doc.add_heading("8.3 Anatomía de una alarma", level=2)
    parrafo(
        doc,
        "Toda regla declara severidad, componente y tres anotaciones en español: qué ocurre, qué "
        "significa **y por qué ese umbral y esa espera**, y qué hacer a continuación.",
    )
    nota(
        doc,
        "**Una alarma sin acción sugerida no está terminada.** Quien la recibe a las tres de la "
        "mañana necesita el siguiente paso, no solo el diagnóstico.",
    )
    doc.add_heading("8.4 Las reglas se prueban", level=2)
    parrafo(
        doc,
        "Las reglas tienen casos de prueba automatizados que se ejecutan en integración continua: "
        "que A1 dispara a los tres minutos y **no** a los dos, que A4 no dispara sin tráfico, que "
        "A3 dispara cuando la serie desaparece. Una regla con un error de sintaxis o mal calibrada "
        "no se descubre hasta que hace falta que dispare —es decir, durante un incidente—, y ese "
        "es el peor momento posible.",
    )
    marcador(doc, "CAPTURA: alarma A1 disparada, con su descripción y su acción", 6.5)
    marcador(doc, "CAPTURA: la misma alarma resuelta tras restablecer el servicio", 6.5)


def seccion_9(doc) -> None:
    doc.add_heading("9. Tableros", level=1)
    tabla(
        doc,
        ["Tablero", "Paneles"],
        [
            ["Salud de la API",
             "Disponibilidad, volumen, latencia p95 (total y por ruta), proporción de errores, "
             "saturación y memoria"],
            ["Negocio",
             "Solicitudes creadas, embudo de cambios de estado, intentos de autenticación por "
             "resultado, altas y respaldos"],
        ],
        [4.0, 11.5],
    )
    parrafo(
        doc,
        "Se aprovisionan de forma declarativa desde archivos versionados, y **se impide editarlos "
        "desde la interfaz**: si se pudieran editar, el archivo versionado dejaría de ser la "
        "verdad en la primera sesión de pruebas.",
    )
    marcador(doc, "CAPTURA: tablero de salud de la API con tráfico real", 8.0)
    marcador(doc, "CAPTURA: tablero de negocio", 8.0)


def seccion_10(doc) -> None:
    doc.add_heading("10. Parámetros de configuración de las herramientas", level=1)

    doc.add_heading("10.1 Colector de OpenTelemetry", level=2)
    tabla(
        doc,
        ["Parámetro", "Valor", "Efecto"],
        [
            ["Distribución", "contrib",
             "El exportador al formato del recolector no viene en la distribución núcleo"],
            ["Receptores", "OTLP por gRPC y HTTP", "Puertos 4317 y 4318"],
            ["Límite de memoria", "Declarado, con margen de pico",
             "Si se queda sin memoria deja de recibir en vez de morir. Un colector muerto "
             "convierte un incidente en dos"],
            ["Agrupación", "Por tiempo y tamaño", "Reduce el número de envíos"],
            ["Atributos de recurso como etiquetas", "Desactivado",
             "Copiarlos a cada serie multiplicaría la cardinalidad sin ganar capacidad de consulta"],
            ["Caducidad de series", "5 minutos",
             "Es el mecanismo que permite detectar la ausencia: sin él, el colector publicaría "
             "indefinidamente el último valor de una aplicación muerta"],
            ["Tubería de trazas", "Creada, con destino provisional",
             "Reservada para el punto 2 del caso de estudio"],
        ],
        [4.5, 4.0, 7.0],
    )

    doc.add_heading("10.2 Recolector de métricas", level=2)
    tabla(
        doc,
        ["Parámetro", "Valor", "Efecto"],
        [
            ["Periodicidad de recolección", "15 s", "Compromiso entre resolución y volumen"],
            ["Periodicidad de evaluación", "15 s", "Las reglas se evalúan al ritmo de los datos"],
            ["Etiquetas externas", "Monitor y ambiente", "Distinguen origen con más de un entorno"],
            ["Archivos de reglas", "Directorio versionado", "Registro y alarmas en archivos distintos"],
            ["Objetivos", "Colector, su telemetría y el sondeo externo",
             "La aplicación no se recoge directamente (sección 2)"],
            ["Plazo del sondeo del servicio publicado", "Ampliado",
             "Un plazo normal reportaría el arranque en frío como caída"],
        ],
        [5.0, 4.0, 6.5],
    )

    doc.add_heading("10.3 Gestor de alertas", level=2)
    tabla(
        doc,
        ["Parámetro", "Valor", "Efecto"],
        [
            ["Agrupación", "Por alarma y ambiente", "Un aviso por hecho, no uno por serie"],
            ["Espera de agrupación", "Menor para las críticas", "Lo urgente sale antes"],
            ["Repetición", "Espaciada", "Evita convertir un incidente largo en ruido continuo"],
            ["Inhibición", "Las derivadas callan si disparó la causa raíz", "Sección 8.2"],
            ["Receptor por omisión", "La propia interfaz", "Sin envíos externos en pruebas"],
            ["Receptor externo", "Documentado y desactivado, contraseña por archivo",
             "Nunca un valor en el archivo de configuración"],
        ],
        [4.0, 5.0, 6.5],
    )

    doc.add_heading("10.4 Tableros", level=2)
    tabla(
        doc,
        ["Parámetro", "Valor", "Efecto"],
        [
            ["Origen de datos", "Aprovisionado y predeterminado", "Cero configuración manual"],
            ["Proveedor de tableros", "Carga desde archivos versionados", "Los tableros son código"],
            ["Edición desde la interfaz", "Desactivada", "El archivo versionado sigue siendo la verdad"],
            ["Registro anónimo y alta de usuarios", "Desactivados", "—"],
            ["Contraseña de administración", "Desde variable de entorno",
             "El arranque se detiene si falta, en lugar de quedarse con la de por defecto"],
            ["Puerto publicado", "3001", "El 3000 lo ocupa el frontend"],
        ],
        [4.5, 4.5, 6.5],
    )

    doc.add_heading("10.5 Sondeo externo", level=2)
    tabla(
        doc,
        ["Parámetro", "Valor", "Efecto"],
        [
            ["Módulo estándar", "Plazo corto", "Para el entorno local"],
            ["Módulo de arranque en frío", "Plazo ampliado", "Para el servicio publicado"],
            ["Códigos aceptados", "Solo 200", "—"],
            ["Afirmación del cuerpo", "Exige el estado sano en la respuesta",
             "No basta el 200: un estado degradado podría devolverlo"],
        ],
        [4.5, 4.5, 6.5],
    )

    doc.add_heading("10.6 Flujos de integración y despliegue", level=2)
    tabla(
        doc,
        ["Parámetro", "Valor", "Efecto"],
        [
            ["Disparador automático", "Imagen verificada sobre la rama de integración",
             "Solo llega lo que ya construyó y arrancó"],
            ["Disparador manual", "Con commit como entrada", "Es también el camino de reversión"],
            ["Puerta de comprobaciones", "Las cuatro, para ese commit", "Sección 4.3"],
            ["Concurrencia", "Agrupada, sin cancelar", "Dos despliegues no migran a la vez"],
            ["Ambiente", "Protegido, con revisor", "Control humano sobre producción"],
            ["Secreto del disparador", "Enmascarado; salida descartada",
             "Nunca aparece en los registros de ejecución"],
            ["Espera inicial", "Antes del primer sondeo", "Sección 4.3"],
            ["Criterio de éxito", "Tres respuestas sanas consecutivas", "Sección 4.3"],
            ["Plazo total", "Construcción más arranque en frío", "Sección 4.3"],
            ["Resumen", "Commits, duración, orden de reversión y advertencia",
             "Quien revierte no improvisa"],
        ],
        [4.5, 5.0, 6.0],
    )

    doc.add_heading("10.7 Fijación de versiones", level=2)
    parrafo(
        doc,
        "Todas las imágenes se fijan a una versión exacta; ninguna etiqueta móvil. Un entorno de "
        "monitoreo que cambia solo es un entorno que un día deja de coincidir con lo documentado, "
        "y el flujo de validación lo comprueba automáticamente.",
    )


def seccion_11(doc) -> None:
    doc.add_heading("11. Guiones del entorno de liberación", level=1)
    tabla(
        doc,
        ["Guion", "Qué hace", "Qué cubre"],
        [
            ["monitoreo-up",
             "Levanta el entorno, espera a que cada herramienta responda sana e imprime sus "
             "direcciones. Se detiene si falta la contraseña del tablero",
             "Generar y configurar el entorno"],
            ["monitoreo-down", "Lo derriba. La purga de datos exige confirmación explícita",
             "Gestionar el entorno"],
            ["generar-trafico",
             "Tráfico representativo: peticiones correctas, rutas inexistentes, inicios de sesión "
             "fallidos, un alta, una solicitud y una transición inválida",
             "Probar en el entorno"],
            ["verificar-monitoreo",
             "Recorre el catálogo completo, comprueba las reglas cargadas y el aprovisionamiento. "
             "Informa qué falta y termina con error",
             "Probar en el entorno"],
            ["desplegar", "Equivalente local del flujo automatizado, con reversión por commit",
             "Desplegar y revertir"],
        ],
        [3.5, 8.0, 4.0],
    )
    parrafo(
        doc,
        "Cada uno en las dos variantes que ya usa el proyecto, sin credenciales dentro y con "
        "código de salida distinto de cero al fallar.",
    )
    nota(
        doc,
        "**El guion de verificación es el mismo que ejecuta la integración continua.** Lo que se "
        "comprueba en un equipo de trabajo es exactamente lo que se comprueba en el pipeline; no "
        "hay dos definiciones de «el monitoreo está bien».",
    )


def seccion_12(doc) -> None:
    doc.add_heading("12. Evidencias", level=1)
    parrafo(
        doc,
        "Todo el trabajo se organizó en dos issues y dos solicitudes de cambio, siguiendo el flujo "
        "del proyecto: un issue por entregable, rama propia con el número del issue, y solicitud "
        "de cambio hacia la rama de integración. Los recuadros amarillos indican **qué captura va "
        "en cada sitio**; el documento no incrusta imágenes, se pegan a mano.",
    )

    doc.add_heading("12.1 Resumen", level=2)
    tabla(
        doc,
        ["Entregable", "Issue", "PR", "Rama", "Commit"],
        [
            ["Documentación: spec 008 y este documento",
             (f"#{dato('documentacion', 'issue')}", dato("documentacion", "issue_url")),
             (f"#{dato('documentacion', 'pr')}", dato("documentacion", "pr_url")),
             dato("documentacion", "rama"), dato("documentacion", "commit")],
            ["Implementación: monitoreo y despliegue",
             (f"#{dato('implementacion', 'issue')}", dato("implementacion", "issue_url")),
             (f"#{dato('implementacion', 'pr')}", dato("implementacion", "pr_url")),
             dato("implementacion", "rama"), dato("implementacion", "commit")],
            ["Issue absorbido: disparador y reversión",
             (f"#{dato('resuelto_189', 'issue')}", dato("resuelto_189", "issue_url")),
             (f"#{dato('implementacion', 'pr')}", dato("implementacion", "pr_url")),
             dato("resuelto_189", "rama"), "—"],
            ["Issue cerrado por cambio de enfoque",
             (f"#{dato('cerrado_187', 'issue')}", dato("cerrado_187", "issue_url")),
             "—", "—", "—"],
        ],
        [5.0, 1.8, 1.8, 5.0, 1.9],
    )

    evidencia(
        doc, "documentacion", "Entregable de documentación",
        [
            f"CAPTURA: issue #{dato('documentacion', 'issue')} en GitHub, con sus etiquetas, "
            "prioridad y estimación",
            f"CAPTURA: PR #{dato('documentacion', 'pr')} — pestaña de conversación, con la "
            "plantilla rellenada y la referencia al issue",
            f"CAPTURA: PR #{dato('documentacion', 'pr')} — pestaña de archivos, mostrando el "
            "árbol de la especificación 008",
            "CAPTURA: comprobaciones de integración continua en verde en el PR",
        ],
    )

    evidencia(
        doc, "implementacion", "Entregable de implementación",
        [
            f"CAPTURA: issue #{dato('implementacion', 'issue')} en GitHub, con sus etiquetas y su "
            "descripción",
            f"CAPTURA: PR #{dato('implementacion', 'pr')} — pestaña de conversación, con "
            "«Resuelve #189» y el enlace a la especificación",
            f"CAPTURA: PR #{dato('implementacion', 'pr')} — pestaña de archivos, con el resumen de "
            "líneas añadidas y quitadas",
            "CAPTURA: comprobaciones de integración continua en verde en el PR",
        ],
    )

    doc.add_heading("Evidencia — Issues resueltos y cerrados", level=2)
    parrafo(
        doc,
        "**#189** se resuelve con la implementación. **#187** se cierra por cambio de enfoque: "
        "pedía el exportador en proceso que la sección 2 descarta, y dejarlo abierto crearía una "
        "contradicción entre el código y el registro de intenciones del proyecto.",
    )
    p = doc.add_paragraph()
    p.add_run("Issue absorbido: ").bold = True
    hipervinculo(p, dato("resuelto_189", "issue_url"), f"#{dato('resuelto_189', 'issue')}")
    p.add_run("   ·   Issue cerrado: ").bold = True
    hipervinculo(p, dato("cerrado_187", "issue_url"), f"#{dato('cerrado_187', 'issue')}")
    marcador(doc, "CAPTURA: issue #189 cerrado, mostrando el PR que lo resolvió", 5.5)
    marcador(
        doc,
        "CAPTURA: issue #187 cerrado, con el comentario que explica el cambio de enfoque y "
        "enlaza la decisión",
        5.5,
    )

    doc.add_heading("12.2 Ejecuciones del pipeline", level=2)
    parrafo(
        doc,
        "Una captura por flujo, como prueba de que cada puerta de la sección 4 existe y pasa.",
    )
    for clave, titulo, descripcion in (
        ("backend_lint", "Formato y análisis",
         "Cero advertencias y cero silenciadores, que es la condición que descarta los paquetes en "
         "preestreno de la sección 2."),
        ("backend_tests", "Pruebas y cobertura",
         "Incluye las pruebas nuevas de las métricas de negocio y la que afirma que la ruta de "
         "métricas responde «no encontrado»."),
        ("frontend_tests", "Compilación y pruebas del frontend",
         "Sin cambios en esta entrega; se incluye porque es una de las cuatro puertas del "
         "despliegue."),
        ("docker_image", "Imagen: construcción y arranque",
         "La imagen arranca contra una base de datos efímera y responde sana."),
        ("monitoring_stack", "Validación del entorno de monitoreo",
         "Configuración de las cinco herramientas, casos de prueba de las reglas de alerta y "
         "verificación de que las métricas del catálogo llegan de verdad."),
        ("deploy", "Despliegue verificado",
         "Puerta de comprobaciones, disparo, espera, sondeo y resumen con la orden de reversión."),
        ("reversion", "Reversión ejecutada",
         "El mismo flujo con el commit sano anterior. Es la evidencia del criterio SLO-9."),
    ):
        ejecucion(doc, clave, titulo, descripcion)

    doc.add_heading("12.3 Mediciones registradas", level=2)
    parrafo(
        doc,
        "Las cifras se toman de ejecuciones reales. Lo que aparezca resaltado en amarillo **no se "
        "ha medido todavía**, y se deja así a propósito.",
    )
    tabla(
        doc,
        ["Medición", "Valor", "Criterio"],
        [
            ["Duración del despliegue, de integrar a servicio sano",
             medicion("duracion_despliegue"), "SLO-10: ≤ 20 min"],
            ["Duración de la reversión", medicion("duracion_reversion"), "SLO-9: ≤ 10 min"],
            ["Métricas del catálogo verificadas", medicion("metricas_verificadas"), "SC-001"],
            ["Reglas de alerta cargadas", medicion("reglas_cargadas"), "10 esperadas"],
            ["Series tras una hora de tráfico", medicion("series_tras_una_hora"),
             "SC-010: por debajo de 2 000"],
            ["Tiempo hasta disparar la alarma de caída", medicion("alerta_tiempo_disparo"),
             "SC-004: ≤ 3 min"],
            ["Tiempo hasta resolverse tras restablecer", medicion("alerta_tiempo_resolucion"),
             "SC-004: ≤ 2 min"],
        ],
        [6.5, 4.5, 4.5],
    )
    marcador(doc, "CAPTURA: salida del guion de verificación del monitoreo, con la tabla completa", 6.0)


def seccion_13(doc) -> None:
    doc.add_heading("13. Punto 2 — Visor de trazabilidad · RESERVADO", level=1)
    nota(
        doc,
        "**No se desarrolla en esta entrega.** Se deja constancia de qué existe ya, para que quien "
        "lo aborde no reconstruya lo hecho.",
    )
    parrafo(doc, "**Ya está en producción:**")
    vinetas(doc, [
        "Registros estructurados en JSON, con redacción automática de secretos y filtrado de las sondas.",
        "Un identificador de traza en **cada línea** de registro, inyectado desde la traza activa.",
        "Instrumentación de trazas de la plataforma web, del cliente HTTP y del controlador de "
        "base de datos, exportándose cuando hay colector.",
        "Una bitácora de negocio en base de datos que incluye el mismo identificador de traza: es "
        "el puente entre lo técnico y lo funcional.",
        "**La tubería de trazas del colector queda creada en esta entrega**, recibiendo y sin "
        "almacén final. Es el punto exacto donde se conecta lo que falta.",
    ])
    parrafo(
        doc,
        "**Lo que falta:** un almacén de trazas y una interfaz de consulta que permita pasar de "
        "una línea de registro a la traza completa de esa petición.",
    )


def seccion_14(doc) -> None:
    doc.add_heading("14. Punto 3 — Visor de auditoría · RESERVADO", level=1)
    nota(doc, "**No se desarrolla en esta entrega.**")
    parrafo(
        doc,
        "**Ya está en producción:** una bitácora de acciones de usuario con su identificador de "
        "traza, y un recurso de consulta en la API.",
    )
    parrafo(
        doc,
        "**Lo que falta:** una interfaz de auditoría con filtros por usuario, acción y periodo; "
        "una política de retención; y la definición de quién puede consultarla y bajo qué registro "
        "de acceso.",
    )


def seccion_15(doc) -> None:
    doc.add_heading("15. Limitaciones y trabajo futuro", level=1)
    parrafo(
        doc,
        "Enumeradas expresamente, porque un documento de monitoreo que no declara sus puntos "
        "ciegos genera más confianza de la que merece.",
    )
    vinetas(doc, [
        "**El entorno de monitoreo es local.** Observa por dentro la aplicación que corre en el "
        "mismo equipo; del servicio publicado solo ve disponibilidad y tiempo de respuesta.",
        "**El propio monitoreo no tiene alta disponibilidad.** Si el equipo se apaga, se deja de "
        "observar. Para un piloto demostrable es aceptable; para operación continua no lo sería.",
        "**Las alarmas no se envían a ningún destino externo.** Hay que entrar a mirarlas.",
        "**La imagen desplegada no es el artefacto verificado**, sino uno equivalente reconstruido "
        "por el proveedor.",
        "**Una migración destructiva no se revierte redesplegando.** La imagen vuelve atrás; los "
        "datos no. La recuperación pasa por los guiones de respaldo, que exigen confirmación.",
        "**No hay separación entre pruebas y producción.** Hoy existe un único ambiente "
        "desplegable; la separación es objeto de la especificación de infraestructura como código.",
        "**El frontend no se monitorea.** El alcance es la aplicación de servidor.",
    ], numeradas=True)


def main() -> None:
    doc = Document()
    configurar(doc)
    portada(doc)
    indice(doc)
    alcance(doc)
    for construir in (
        seccion_1, seccion_2, seccion_3, seccion_4, seccion_5, seccion_6, seccion_7,
        seccion_8, seccion_9, seccion_10, seccion_11, seccion_12, seccion_13,
        seccion_14, seccion_15,
    ):
        doc.add_page_break()
        construir(doc)
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    doc.save(SALIDA)
    print(f"Generado: {SALIDA}")
    pendientes = sum(
        1 for item in ITEMS.values() for valor in item.values() if valor in (None, "")
    ) + sum(1 for valor in MEDICIONES.values() if valor in (None, ""))
    if pendientes:
        print(f"Quedan {pendientes} datos pendientes, resaltados en amarillo en el documento.")


if __name__ == "__main__":
    main()
